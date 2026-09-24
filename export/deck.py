import logging
import math
from io import BytesIO

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

log = logging.getLogger(__name__)
MAX_TITLE_WORDS = 12
BODY_WORD_BUDGET = 40
MAX_BULLETS = 3
TABLE_ROWS, TABLE_COLS = 8, 5
CHART_TYPES = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
               "pie": XL_CHART_TYPE.PIE, "scatter": XL_CHART_TYPE.XY_SCATTER}


def truncate_title(text: str, max_words: int = MAX_TITLE_WORDS) -> str:
    words = text.split()
    return text if len(words) <= max_words else " ".join(words[:max_words]) + "…"


def _cut(text: str, words: int) -> str:
    parts = text.split()
    return text if len(parts) <= words else " ".join(parts[:words])


def fit_body(evidence: list[str], recommendation: str, budget: int = BODY_WORD_BUDGET) -> tuple[list[str], str]:
    rec = _cut(recommendation, budget)
    remaining = budget - len(rec.split())
    bullets: list[str] = []
    for item in evidence[:MAX_BULLETS]:
        if remaining <= 0:
            break
        cut = _cut(item, remaining)
        bullets.append(cut)
        remaining -= len(cut.split())
    return bullets, rec


def _clean(text) -> str:
    return "".join(ch for ch in str(text) if ch >= " " or ch in "\n\t")


def _cell_text(value) -> str:
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    return _clean(f"{value:.4g}" if isinstance(value, float) else value)


def _add_table(slide, df: pd.DataFrame) -> None:
    df = df.iloc[:TABLE_ROWS, :TABLE_COLS]
    shape = slide.shapes.add_table(len(df) + 1, len(df.columns), Inches(6.6), Inches(1.7), Inches(6.2), Inches(0.4) * (len(df) + 1))
    table = shape.table
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = _clean(col)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            table.cell(i, j).text = _cell_text(value)


def _build_chart(slide, chart: dict, df: pd.DataFrame) -> bool:
    x, y, kind = chart.get("x"), chart.get("y"), chart.get("type")
    if kind not in CHART_TYPES or x not in df.columns or y not in df.columns:
        return False
    df = df.assign(**{y: pd.to_numeric(df[y], errors="coerce")})
    if kind == "scatter":
        df = df.assign(**{x: pd.to_numeric(df[x], errors="coerce")})
    df = df.dropna(subset=[x, y])
    df = df[df[y].map(math.isfinite)]
    if kind == "scatter":
        df = df[df[x].map(math.isfinite)]
    if df.empty:
        return False
    title = _clean(chart.get("title") or "")
    if kind == "scatter":
        data = XyChartData()
        series = data.add_series(title or str(y))
        for xv, yv in zip(df[x], df[y]):
            series.add_data_point(float(xv), float(yv))
    else:
        data = CategoryChartData()
        data.categories = [_clean(v) for v in df[x]]
        data.add_series(_clean(y), [float(v) for v in df[y]])
    frame = slide.shapes.add_chart(CHART_TYPES[kind], Inches(6.6), Inches(1.7), Inches(6.2), Inches(5), data)
    if title:
        frame.chart.has_title = True
        frame.chart.chart_title.text_frame.text = title
    return True


def _add_chart(slide, chart, df: pd.DataFrame) -> bool:
    if not isinstance(chart, dict):
        return False
    n = len(slide.shapes)
    try:
        return _build_chart(slide, chart, df)
    except Exception:
        log.debug("chart creation failed; falling back to table", exc_info=True)
        for shape in list(slide.shapes)[n:]:  # remove any partially added shape
            shape._element.getparent().remove(shape._element)
        return False


def build_deck(entries: list[dict]) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for e in entries:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        ins = e["insight"]
        title = slide.shapes.title
        title.text = _clean(truncate_title(ins["finding"]))
        title.left, title.top, title.width, title.height = Inches(0.6), Inches(0.3), Inches(12.1), Inches(1.25)
        bullets, rec = fit_body(ins["evidence"], ins["recommendation"])
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(5.7), Inches(5))
        frame = box.text_frame
        frame.word_wrap = True
        for i, text in enumerate(bullets):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = _clean(f"• {text}")
            p.font.size = Pt(18)
        p = frame.paragraphs[0] if not bullets else frame.add_paragraph()
        p.text = _clean(f"Recommendation: {rec}")
        p.font.size = Pt(18)
        p.font.bold = True
        df = pd.DataFrame(e.get("slice") or [])
        if df.empty:
            continue
        if not (e.get("chart") and _add_chart(slide, e["chart"], df)):
            _add_table(slide, df)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
