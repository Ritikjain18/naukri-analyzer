from io import BytesIO

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from export.deck import build_deck, fit_body, truncate_title


def entry(chart_type="bar", finding="Engineering has the highest conversion rate.", chart=True, slice_rows=None):
    rows = slice_rows if slice_rows is not None else [{"category": "Eng", "rate": 0.2}, {"category": "Sales", "rate": 0.1}]
    return {"question": "q", "approved": False,
            "insight": {"finding": finding, "evidence": ["Eng 20%", "Sales 10%", "Gap 10 points", "Extra bullet"],
                        "recommendation": "Shift budget to Engineering."},
            "chart": {"type": chart_type, "x": "category", "y": "rate", "title": "Conversion"} if chart else None,
            "slice": rows}


def load(entries):
    return Presentation(BytesIO(build_deck(entries)))


def shapes_of(slide):
    return list(slide.shapes)


def test_truncate_title():
    assert truncate_title("one two three") == "one two three"
    long = " ".join(f"w{i}" for i in range(20))
    out = truncate_title(long)
    assert out.endswith("…") and len(out.rstrip("…").split()) == 12


def test_fit_body_limits_words_and_bullets():
    ev = ["word " * 20, "two", "three", "four"]
    bullets, rec = fit_body(ev, "Do the thing now.", budget=40)
    assert len(bullets) <= 3 and rec.startswith("Do the thing")
    assert sum(len(b.split()) for b in bullets) + len(rec.split()) <= 40


def test_fit_body_cuts_recommendation_when_alone_too_long():
    bullets, rec = fit_body(["x"], "word " * 60, budget=40)
    assert len(rec.split()) <= 40 and bullets == []


def test_one_slide_per_entry_with_title():
    prs = load([entry(), entry(finding="Sales is lowest.")])
    assert len(prs.slides) == 2
    assert prs.slides[0].shapes.title.text == "Engineering has the highest conversion rate."


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_native_chart_present(kind):
    slide = load([entry(chart_type=kind)]).slides[0]
    expected = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[kind]
    charts = [s.chart for s in shapes_of(slide) if s.has_chart]
    assert charts and charts[0].chart_type == expected


def test_scatter_with_numeric_x_and_fallback_to_table():
    rows = [{"category": 1.0, "rate": 0.2}, {"category": 2.0, "rate": 0.1}]
    assert any(s.has_chart for s in shapes_of(load([entry("scatter", slice_rows=rows)]).slides[0]))
    assert any(s.has_table for s in shapes_of(load([entry("scatter")]).slides[0]))   # x non-numeric


def test_table_fallback_when_no_chart():
    slide = load([entry(chart=False)]).slides[0]
    assert any(s.has_table for s in shapes_of(slide)) and not any(s.has_chart for s in shapes_of(slide))
    table = [s for s in shapes_of(slide) if s.has_table][0].table
    assert [table.cell(0, j).text for j in range(2)] == ["category", "rate"]


def test_no_visual_when_slice_empty():
    slide = load([entry(chart=False, slice_rows=[])]).slides[0]
    assert not any(s.has_table or s.has_chart for s in shapes_of(slide))


def test_body_text_contains_recommendation_and_bullet_limit():
    slide = load([entry()]).slides[0]
    text = " ".join(s.text_frame.text for s in shapes_of(slide) if s.has_text_frame)
    assert "Recommendation: Shift budget to Engineering." in text and "Extra bullet" not in text


def kinds(slide):
    return {"chart": any(s.has_chart for s in shapes_of(slide)), "table": any(s.has_table for s in shapes_of(slide))}


def test_title_geometry_spans_slide():
    t = load([entry()]).slides[0].shapes.title
    assert t.left == Inches(0.6) and t.width == Inches(12.1) and t.top == Inches(0.3) and t.height == Inches(1.25)


NAN = float("nan")


@pytest.mark.parametrize("kind", ["bar", "line", "pie", "scatter"])
def test_nan_and_none_in_y_do_not_crash(kind):
    rows = [{"category": 1.0, "rate": 0.2}, {"category": 2.0, "rate": NAN}, {"category": 3.0, "rate": None}]
    slide = load([entry(kind, slice_rows=rows)]).slides[0]
    assert kinds(slide)["chart"]


@pytest.mark.parametrize("kind", ["bar", "scatter"])
def test_all_nan_y_falls_back_to_table(kind):
    rows = [{"category": 1.0, "rate": NAN}, {"category": 2.0, "rate": None}]
    k = kinds(load([entry(kind, slice_rows=rows)]).slides[0])
    assert k["table"] and not k["chart"]


def test_table_blanks_nan_and_none():
    rows = [{"category": "Eng", "rate": NAN}, {"category": None, "rate": 0.1}]
    slide = load([entry(chart=False, slice_rows=rows)]).slides[0]
    table = [s for s in shapes_of(slide) if s.has_table][0].table
    assert table.cell(1, 1).text == "" and table.cell(2, 0).text == ""


@pytest.mark.parametrize("chart", [
    {"type": "bar", "x": "category", "y": "rate", "title": None},
    {"type": "bar", "x": "category", "title": "t"},
    {"type": "donut", "x": "category", "y": "rate", "title": "t"},
    {"x": "category", "y": "rate"},
])
def test_malformed_chart_dicts_degrade(chart):
    e = entry()
    e["chart"] = chart
    prs = load([e, entry()])
    assert len(prs.slides) == 2
    if chart.get("type") == "bar" and chart.get("y"):
        assert kinds(prs.slides[0])["chart"]
    else:
        assert kinds(prs.slides[0])["table"]


def test_control_characters_stripped():
    e = entry(finding="Bad\x0b title\x00 here")
    e["insight"]["evidence"] = ["ev\x0bidence"]
    e["slice"] = [{"category": "a\x00b", "rate": 1}]
    e["chart"] = None
    slide = load([e]).slides[0]
    assert slide.shapes.title.text == "Bad title here"
    text = " ".join(s.text_frame.text for s in shapes_of(slide) if s.has_text_frame)
    assert "\x0b" not in text and "\x00" not in text
    table = [s for s in shapes_of(slide) if s.has_table][0].table
    assert table.cell(1, 0).text == "ab"


def test_chart_creation_exception_falls_back_to_table(monkeypatch):
    import export.deck as deck
    monkeypatch.setattr(deck, "_build_chart", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    assert kinds(load([entry()]).slides[0])["table"]


def test_deck_survives_xml_forbidden_characters():
    from io import BytesIO

    from pptx import Presentation

    from export.deck import build_deck

    bad = "a￾b￿c\ud800d\x00e\x0bf"
    entry = {"question": "q", "approved": True,
             "insight": {"finding": bad, "evidence": [bad, 5], "recommendation": bad},
             "slice": [{bad: bad, "n": 1}, {bad: "z", "n": 2}],
             "chart": {"type": "bar", "x": bad, "y": "n", "title": bad}}
    table_entry = {**entry, "chart": None}
    prs = Presentation(BytesIO(build_deck([entry, table_entry])))
    assert len(prs.slides) == 2
    assert "abcdef" in prs.slides[0].shapes.title.text


def test_entry_label_tolerates_missing_finding():
    from export.deck import entry_label

    assert entry_label(0, {"insight": {"finding": None}}) == "1. (no finding)"
    assert entry_label(1, {}) == "2. (no finding)"
